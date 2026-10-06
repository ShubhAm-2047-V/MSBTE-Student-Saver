import docx

doc = docx.Document('ITR.report...docx')
print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))

for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if 'w:br' in xml and 'type="page"' in xml:
        print(f"PageBreak at para {i}: text='{p.text[:40]}'")
    elif p.style.name.startswith('Heading 1'):
        print(f"Heading 1 at para {i}: '{p.text}'")
