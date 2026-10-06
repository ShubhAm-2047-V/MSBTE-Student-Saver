import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_floating_background(paragraph, image_path, width_inches=7.0):
    """Inserts an image as a floating background behind text (watermark/backdrop)."""
    run = paragraph.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    
    drawings = run._r.xpath('.//w:drawing')
    if not drawings:
        return
    drawing = drawings[0]
    inlines = drawing.xpath('.//wp:inline')
    if not inlines:
        return
    inline = inlines[0]
    
    graphic = inline.xpath('.//a:graphic')[0]
    extent = inline.xpath('.//wp:extent')[0]
    cx = extent.get('cx')
    cy = extent.get('cy')
    
    anchor_xml = parse_xml(
        f'<wp:anchor {nsdecls("wp")} {nsdecls("a")} {nsdecls("pic")} '
        f'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="101" name="PageBackgroundArt"/>'
        f'<wp:cNvGraphicFramePr/>'
        f'</wp:anchor>'
    )
    anchor_xml.append(graphic)
    drawing.remove(inline)
    drawing.append(anchor_xml)

doc = Document()
p1 = doc.add_paragraph()
add_floating_background(p1, 'static/images/login_bg_watermark.jpg', width_inches=7.2)
doc.add_paragraph("This is text over the background image on page 1")
doc.save("test_watermark_out.docx")
print("Successfully generated test_watermark_out.docx")
