import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
sec = doc.sections[0]
sec.different_first_page_header_footer = True

footer = sec.footer
p = footer.paragraphs[0]
p.text = "" # clear

pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="double" w:sz="12" w:space="6" w:color="A67C52"/></w:pBdr>')
p._p.get_or_add_pPr().append(pBdr)

tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="9792"/></w:tabs>')
p._p.get_or_add_pPr().append(tabs)

r1 = p.add_run("MSBTE Student Saver\t")
r1.bold = True
r1.italic = True
r1.font.name = "Times New Roman"
r1.font.size = Pt(9.5)
r1.font.color.rgb = RGBColor(100, 116, 139)

fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="19"/><w:color w:val="64748B"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
p._p.append(fld)

doc.add_paragraph("Cover Page Content")
doc.add_page_break()
doc.add_paragraph("Page 2 Content with bottom border and app name footer")
doc.save("test_footer_out.docx")
print("Saved test footer successfully!")
